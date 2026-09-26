#!/usr/bin/env python3
import sys
from openpyxl import load_workbook
# Sheet contents include non-ASCII (pt-BR, "Δ"); Windows consoles default to cp1252.
sys.stdout.reconfigure(encoding="utf-8")
wb = load_workbook(sys.argv[1])
for ws in wb.worksheets:
    print(f"== {ws.title} ({ws.max_row}x{ws.max_column})")
    for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, 3), values_only=True):
        print("   ", row)
print("OK", len(wb.worksheets), "sheets")
